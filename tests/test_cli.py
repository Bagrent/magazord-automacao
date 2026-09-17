"""End-to-end runs through the command line, against a mocked API."""

from __future__ import annotations

import httpx
import pytest
import respx
from openpyxl import Workbook, load_workbook

from magazord_automacao.cli import main


def planilha(tmp_path):
    wb = Workbook()
    wb.active.append(["SKU", "Nome", "Preço", "Estoque", "Imagem 1"])
    wb.active.append(["A1", "Camiseta", "89,90", 10, "https://cdn/1.jpg"])
    wb.active.append(["A2", "Boné", "49,90", 5, "https://cdn/2.jpg"])
    caminho = tmp_path / "produtos.xlsx"
    wb.save(caminho)
    return caminho


def test_validar_aprova_uma_planilha_boa(tmp_path, capsys):
    assert main(["validar", str(planilha(tmp_path))]) == 0
    assert "2 produto(s) válido(s)" in capsys.readouterr().out


def test_validar_sinaliza_linha_ruim(tmp_path, capsys):
    wb = Workbook()
    wb.active.append(["SKU", "Nome", "Preço"])
    wb.active.append(["A1", "Camiseta", "muito caro"])
    caminho = tmp_path / "ruim.xlsx"
    wb.save(caminho)

    assert main(["validar", str(caminho)]) == 1
    assert "rejeitada" in capsys.readouterr().out


def test_enviar_simula_por_padrao(tmp_path, api_config, base, capsys, monkeypatch):
    monkeypatch.setenv("MAGAZORD_BASE_URL", base)
    relatorio = tmp_path / "rel.xlsx"

    with respx.mock(assert_all_called=False) as mock:
        criar = mock.post(f"{base}/site/produto").mock(return_value=httpx.Response(201))
        codigo = main([
            "enviar", str(planilha(tmp_path)),
            "--sem-imagens",
            "--estado", str(tmp_path / "e.sqlite3"),
            "--relatorio", str(relatorio),
        ])

    assert codigo == 0
    assert not criar.called                      # nothing sent without --confirmar
    assert "SIMULAÇÃO" in capsys.readouterr().out
    assert relatorio.exists()


def test_enviar_com_confirmar_cria_os_produtos(tmp_path, api_config, base, monkeypatch):
    monkeypatch.setenv("MAGAZORD_BASE_URL", base)
    monkeypatch.setenv("MAGAZORD_USER", "u")
    monkeypatch.setenv("MAGAZORD_TOKEN", "t")
    relatorio = tmp_path / "rel.xlsx"

    with respx.mock(assert_all_called=False) as mock:
        mock.get(f"{base}/site/produto").mock(
            return_value=httpx.Response(200, json={"items": []})
        )
        criar = mock.post(f"{base}/site/produto").mock(
            return_value=httpx.Response(201, json={"id": 9})
        )
        codigo = main([
            "enviar", str(planilha(tmp_path)), "--confirmar", "--sem-imagens",
            "--estado", str(tmp_path / "e.sqlite3"), "--relatorio", str(relatorio),
        ])

    assert codigo == 0
    assert criar.call_count == 2

    ws = load_workbook(relatorio)["Envio"]
    linhas = list(ws.iter_rows(min_row=2, values_only=True))
    assert [l[1] for l in linhas] == ["A1", "A2"]
    assert {l[3] for l in linhas} == {"criado"}


def test_enviar_sem_credenciais_para_antes(tmp_path, api_config, base, monkeypatch, capsys):
    monkeypatch.setenv("MAGAZORD_BASE_URL", base)
    monkeypatch.delenv("MAGAZORD_USER", raising=False)
    monkeypatch.delenv("MAGAZORD_TOKEN", raising=False)

    codigo = main(["enviar", str(planilha(tmp_path)), "--confirmar",
                   "--estado", str(tmp_path / "e.sqlite3")])
    assert codigo == 2
    assert "credenciais ausentes" in capsys.readouterr().err


# -- catálogo ---------------------------------------------------------------

def _ofertas(tmp_path):
    wb = Workbook()
    wb.active.append(["Produto", "Preço de", "Preço por", "Imagem"])
    wb.active.append(["Base Real Filter", "59,90", "39,90", ""])
    wb.active.append(["Gloss Chocochilli", "59,90", "29,90", ""])
    caminho = tmp_path / "ofertas.xlsx"
    wb.save(caminho)
    return caminho


def test_catalogo_monta_flipbook_e_pdf(tmp_path, capsys):
    pytest.importorskip("PIL", reason="o catálogo depende do Pillow")
    saida = tmp_path / "catalogo"

    codigo = main(["catalogo", str(_ofertas(tmp_path)), "--saida", str(saida),
                   "--titulo", "Promoções de Setembro"])

    assert codigo == 0
    assert (saida / "index.html").exists()
    assert (saida / "catalogo.pdf").exists()
    assert len(list((saida / "paginas").glob("*.jpg"))) == 3      # capa + 2 ofertas
    saida_texto = capsys.readouterr().out
    assert "2 oferta(s)" in saida_texto and "3 página(s)" in saida_texto


def test_catalogo_aceita_uma_pasta_de_artes_prontas(tmp_path):
    pytest.importorskip("PIL", reason="o catálogo depende do Pillow")
    from PIL import Image

    artes = tmp_path / "artes"
    artes.mkdir()
    for i in (2, 1):                                   # fora de ordem de propósito
        Image.new("RGB", (216, 384), (40 * i, 30, 20)).save(artes / f"{i:02d}.png")

    assert main(["catalogo", str(artes), "--saida", str(tmp_path / "c"), "--sem-pdf"]) == 0
    assert not (tmp_path / "c" / "catalogo.pdf").exists()
    # A arte pronta entra como está: PNG continua PNG, sem passar pelo JPEG.
    paginas = sorted((tmp_path / "c" / "paginas").iterdir())
    assert [p.name for p in paginas] == ["pagina-01.png", "pagina-02.png"]
    assert paginas[0].read_bytes() == (artes / "01.png").read_bytes()


def test_catalogo_pode_recodificar_quando_pedido(tmp_path):
    pytest.importorskip("PIL", reason="o catálogo depende do Pillow")
    from PIL import Image

    artes = tmp_path / "artes"
    artes.mkdir()
    Image.new("RGB", (216, 384), (200, 120, 30)).save(artes / "01.png")

    assert main(["catalogo", str(artes), "--saida", str(tmp_path / "c"),
                 "--sem-pdf", "--recodificar"]) == 0
    assert [p.name for p in (tmp_path / "c" / "paginas").iterdir()] == ["pagina-01.jpg"]


def test_catalogo_avisa_quando_a_entrada_nao_existe(tmp_path, capsys):
    assert main(["catalogo", str(tmp_path / "nao-existe.xlsx")]) == 2
    assert "não encontrei" in capsys.readouterr().err


def test_modelo_de_ofertas_sai_no_formato_do_catalogo(tmp_path, capsys):
    destino = tmp_path / "modelo.xlsx"
    assert main(["modelo", "--tipo", "ofertas", "--destino", str(destino)]) == 0
    assert "Produto" in [c.value for c in load_workbook(destino).active[1]]


def test_catalogo_recusa_gravar_dentro_da_pasta_das_artes(tmp_path, capsys):
    pytest.importorskip("PIL", reason="o catálogo depende do Pillow")
    from PIL import Image

    artes = tmp_path / "artes"
    artes.mkdir()
    Image.new("RGB", (216, 384), (200, 120, 30)).save(artes / "01.png")

    assert main(["catalogo", str(artes), "--saida", str(artes / "catalogo")]) == 2
    assert "dentro da pasta das artes" in capsys.readouterr().err
    assert not (artes / "catalogo").exists()

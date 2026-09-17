"""The example offers sheet -- the format ``magazord catalogo`` reads.

Filling a spreadsheet is the part of this that a shop assistant can do, so the
file documents itself: real-looking rows and a cell comment on every header
explaining what belongs in that column. The explanations are comments rather
than a second header row on purpose -- a help row would be read back as a
product and printed as a page.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUNAS = [
    ("Produto", 44, "nome que aparece na página"),
    ("Preço de", 14, "preço antigo; deixe vazio para não mostrar o 'DE:'"),
    ("Preço por", 14, "preço da promoção"),
    ("Imagem", 46, "arquivo na pasta ao lado da planilha, ou link https://"),
    ("Selo", 14, "opcional: NOVIDADE, ÚLTIMAS UNIDADES..."),
    ("Link", 46, "opcional: leva o botão COMPRAR ao produto na loja"),
    ("Observação", 34, "opcional: linha extra no rodapé da página"),
    ("SKU", 18, "opcional: só para você se localizar"),
]

LINHAS = [
    ["Base Liquida Franciny Ehlke Real Filter", "59,99", "39,99",
     "fotos/base-real-filter.png", "", "https://loja.exemplo.com.br/base-real-filter", "", "FE-BASE-01"],
    ["Gloss Franciny Ehlke Chocochilli", "59,99", "29,99",
     "fotos/gloss-chocochilli.png", "", "", "", "FE-GLOSS-02"],
    ["Batom Liquido Franciny Ehlke", "41,99", "29,99",
     "fotos/batom-liquido.png", "4 cores", "", "", "FE-BATOM-03"],
    ["Gloss Franciny Ehlke Franboesa 5g", "69,99", "39,99",
     "fotos/gloss-franboesa.png", "NOVIDADE", "", "", "FE-GLOSS-04"],
    ["Gloss Franciny Ehlke Glossip Girl 4,5ml", "42,99", "39,99",
     "fotos/glossip-girl.png", "", "", "", "FE-GLOSS-05"],
]


def gerar(destino: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ofertas"

    cabecalho = Font(bold=True, color="FFFFFF")
    fundo = PatternFill("solid", fgColor="EC7A18")

    for i, (titulo, largura, explicacao) in enumerate(COLUNAS, start=1):
        celula = ws.cell(row=1, column=i, value=titulo)
        celula.font = cabecalho
        celula.fill = fundo
        celula.alignment = Alignment(horizontal="center", vertical="center")
        celula.comment = Comment(explicacao, "magazord")
        ws.column_dimensions[get_column_letter(i)].width = largura

    for linha in LINHAS:
        ws.append(linha)

    ws.freeze_panes = "A2"
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino

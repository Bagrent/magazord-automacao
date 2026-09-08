"""Writing the per-run report the operator actually reads."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .pipeline import Resumo

CORES = {
    "criado": "C6EFCE",
    "atualizado": "DDEBF7",
    "inalterado": "F2F2F2",
    "simulado": "FFF2CC",
    "erro": "FFC7CE",
}

COLUNAS_ENVIO = [
    ("Linha", 8), ("SKU", 20), ("Nome", 40), ("Status", 14),
    ("ID Magazord", 16), ("Imagens", 10), ("Mensagem", 70),
]
COLUNAS_ERRO = [("Linha", 8), ("SKU", 20), ("Coluna", 20), ("Problema", 80)]


def _cabecalho(ws, colunas) -> None:
    fonte = Font(bold=True, color="FFFFFF")
    fundo = PatternFill("solid", fgColor="2F5496")
    for i, (titulo, largura) in enumerate(colunas, start=1):
        celula = ws.cell(row=1, column=i, value=titulo)
        celula.font = fonte
        celula.fill = fundo
        celula.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.freeze_panes = "A2"


def gerar(resumo: Resumo, destino: Path) -> Path:
    """Write an .xlsx with one tab per concern: the run, and the rejected rows."""
    wb = Workbook()

    ws = wb.active
    ws.title = "Envio"
    _cabecalho(ws, COLUNAS_ENVIO)
    for r in resumo.resultados:
        ws.append([r.linha, r.sku, r.nome, r.status, r.id_magazord or "",
                   r.imagens_enviadas, r.mensagem])
        cor = CORES.get(r.status)
        if cor:
            ws.cell(row=ws.max_row, column=4).fill = PatternFill("solid", fgColor=cor)

    if resumo.erros_leitura:
        wse = wb.create_sheet("Linhas rejeitadas")
        _cabecalho(wse, COLUNAS_ERRO)
        for e in resumo.erros_leitura:
            wse.append([e.linha, e.sku or "", e.campo or "", e.mensagem])
            wse.cell(row=wse.max_row, column=4).fill = PatternFill("solid", fgColor=CORES["erro"])

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino


def caminho_padrao(pasta: Path) -> Path:
    carimbo = datetime.now().strftime("%Y%m%d-%H%M%S")
    return pasta / f"relatorio-{carimbo}.xlsx"


def resumo_texto(resumo: Resumo) -> str:
    linhas = [
        f"  criados .......... {resumo.criados}",
        f"  atualizados ...... {resumo.atualizados}",
        f"  inalterados ...... {resumo.inalterados}",
    ]
    if resumo.simulados:
        linhas.append(f"  simulados ........ {resumo.simulados}")
    linhas += [
        f"  erros de envio ... {resumo.erros}",
        f"  linhas rejeitadas  {len(resumo.erros_leitura)}",
        f"  imagens enviadas . {resumo.imagens}",
    ]
    return "\n".join(linhas)

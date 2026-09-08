"""Generation of the example workbook that documents the expected input format."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUNAS = [
    ("SKU", 16), ("Nome", 38), ("Descrição", 46), ("EAN", 16), ("Marca", 16),
    ("Categoria", 26), ("Preço", 12), ("Preço Promocional", 18), ("Custo", 12),
    ("Estoque", 10), ("Peso", 10), ("Altura", 10), ("Largura", 10),
    ("Comprimento", 13), ("NCM", 12), ("Ativo", 8), ("SKU Pai", 14),
    ("Cor", 14), ("Tamanho", 10),
    ("Imagem 1", 40), ("Imagem 2", 40), ("Imagem 3", 40),
]

LINHAS = [
    ["CAM-001-P-AZ", "Camiseta Básica Algodão Pima", "Camiseta de algodão pima 30.1, gola careca, corte regular.",
     "7891234567895", "Bagrent", "Roupas > Camisetas", "89,90", "79,90", "32,00",
     25, "0,180", 2, 20, 30, "6109.10.00", "Sim", "CAM-001", "Azul", "P",
     "https://cdn.exemplo.com.br/cam-001-azul-frente.jpg",
     "https://cdn.exemplo.com.br/cam-001-azul-costas.jpg", ""],
    ["CAM-001-M-AZ", "Camiseta Básica Algodão Pima", "Camiseta de algodão pima 30.1, gola careca, corte regular.",
     "7891234567901", "Bagrent", "Roupas > Camisetas", "89,90", "79,90", "32,00",
     40, "0,190", 2, 21, 31, "6109.10.00", "Sim", "CAM-001", "Azul", "M",
     "https://cdn.exemplo.com.br/cam-001-azul-frente.jpg",
     "https://cdn.exemplo.com.br/cam-001-azul-costas.jpg", ""],
    ["TEN-455-42", "Tênis Runner Trail 455", "Tênis de corrida em trilha, solado EVA, cabedal em mesh.",
     "7899876543215", "Bagrent", "Calçados > Tênis", "349,90", "", "150,00",
     12, "0,820", 12, 22, 33, "6404.11.00", "Sim", "", "Preto", "42",
     "https://cdn.exemplo.com.br/ten-455-lateral.jpg",
     "https://cdn.exemplo.com.br/ten-455-superior.jpg",
     "https://cdn.exemplo.com.br/ten-455-solado.jpg"],
]


def gerar(destino: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Produtos"

    cabecalho = Font(bold=True, color="FFFFFF")
    fundo = PatternFill("solid", fgColor="2F5496")
    for i, (titulo, largura) in enumerate(COLUNAS, start=1):
        celula = ws.cell(row=1, column=i, value=titulo)
        celula.font = cabecalho
        celula.fill = fundo
        celula.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = largura

    for linha in LINHAS:
        ws.append(linha)

    ws.freeze_panes = "A2"
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino
